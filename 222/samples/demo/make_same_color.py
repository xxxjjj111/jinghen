# -*- coding: utf-8 -*-
"""生成「同色隐藏文字」演示样本（⑦）：白字白底的 docx，肉眼看不见，全选/复制能带出敏感字"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_COLOR_INDEX

doc = Document()
doc.add_heading("关于2026年度合作项目报价的说明", level=1)

doc.add_paragraph("一、项目概况")
doc.add_paragraph("本项目为杭州XX科技有限公司的外包开发合同，合同金额及商务条款如下。")

# 正常可见内容
doc.add_paragraph("二、商务条款（对外版）")
p1 = doc.add_paragraph("合同金额：人民币120万元整；付款分三期。")
r1 = p1.add_run("（以上为对外口径，具体以合同为准）")
r1.font.color.rgb = RGBColor(0, 0, 0)

# 同色隐藏：白字（在白色页面上肉眼不可见，但全选/复制能带出）
doc.add_paragraph("三、内部附注（此段为白字，页面上不可见，全选可复制）")
p2 = doc.add_paragraph()
r2 = p2.add_run("内部底价：86万元；实际回扣率30%；客户联系人张教授；"
                "最低可接受折扣至7.5折；内部预算密码998877。")
r2.font.color.rgb = RGBColor(255, 255, 255)  # 白色字体，白底 → 同色隐藏

# 再一段正常内容
doc.add_paragraph("四、联系方式")
doc.add_paragraph("联系人：李同学（市场部），电话0571-8888****。")

out = "/home/user/Doubao/chats/38447830452557314/_wang2/网安第二版/samples/demo/演示7_同色隐藏文字.docx"
doc.save(out)
print("写入", out)
